"""Provenance audit test for all 6 generated Excel workbooks.
Validates Decision JP-D03 in decision campaign cil-live-job-provenance-20261018.
"""
from __future__ import annotations

from pathlib import Path
import openpyxl
import pytest

PROJECT_ROOT = Path('/home/ubuntu/root/company-intelligence-platform')

WORKBOOKS = [
    (PROJECT_ROOT / 'korea-public-sector-it-management-jobs.xlsx', 423),
    (PROJECT_ROOT / 'czech-public-sector-it-management-jobs.xlsx', 175),
    (PROJECT_ROOT / 'slovakia-public-sector-it-management-jobs.xlsx', 10),
    (PROJECT_ROOT / 'korea-graduate-rd-it-management-jobs.xlsx', 107),
    (PROJECT_ROOT / 'czech-prague-graduate-rd-it-management-jobs.xlsx', 21),
    (PROJECT_ROOT / 'slovakia-bratislava-graduate-rd-it-management-jobs.xlsx', 5),
]


@pytest.mark.parametrize("file_path, expected_rows", WORKBOOKS)
def test_workbook_provenance_and_hyperlinks(file_path: Path, expected_rows: int):
    assert file_path.exists(), f"Workbook {file_path.name} does not exist"

    wb = openpyxl.load_workbook(file_path, data_only=False)
    assert "메타데이터_및_검증정보" in wb.sheetnames

    total_data_rows = 0
    valid_url_count = 0
    valid_hyperlink_formula_count = 0
    valid_freshness_count = 0

    for sheet_name in wb.sheetnames:
        if sheet_name == "메타데이터_및_검증정보":
            continue

        ws = wb[sheet_name]
        # Verify 30 header columns at row 4
        assert ws.max_column >= 30, f"Sheet {sheet_name} in {file_path.name} has only {ws.max_column} columns, expected 30"

        # Check headers
        header_vals = [ws.cell(row=4, column=col).value for col in range(1, 31)]
        assert header_vals[4] == "법인 형태 (Legal Form)"  # Col 5
        assert header_vals[5] == "공공 분류 / 기관 유형 (Entity Type)"  # Col 6
        assert header_vals[14] == "주간 근무시간 / 모집인원 (Work Hours / Headcount)"  # Col 15
        assert header_vals[17] == "통화 (Currency)"  # Col 18
        assert header_vals[18] == "근무지 / 소재지 (Location)"  # Col 19
        assert header_vals[19] == "원천 공고 URL (Posting URL)"  # Col 20
        assert header_vals[20] == "공고 바로가기 (Clickable Hyperlink)"  # Col 21
        assert header_vals[21] == "공고 실시간 상태 (Listing Freshness)"  # Col 22
        assert header_vals[28] == "무결성 검증 (Join Verified)"  # Col 29
        assert header_vals[29] == "검증 메모 (Verification Note)"  # Col 30

        # Iterate rows
        for r_idx in range(5, ws.max_row + 1):
            val_no = ws.cell(row=r_idx, column=1).value
            if val_no is None:
                continue

            total_data_rows += 1
            raw_url = str(ws.cell(row=r_idx, column=20).value or "").strip()
            hyperlink_formula = str(ws.cell(row=r_idx, column=21).value or "").strip()
            freshness = str(ws.cell(row=r_idx, column=22).value or "").strip()

            # 1. URL must not be empty or 'None'
            assert raw_url and raw_url.lower() != "none" and raw_url.startswith("http"), (
                f"Row {r_idx} in {sheet_name} ({file_path.name}) has invalid URL: {raw_url}"
            )
            valid_url_count += 1

            # 2. Hyperlink formula must match =HYPERLINK(T{row}, "공고 확인 ↗")
            expected_formula = f'=HYPERLINK(T{r_idx}, "공고 확인 ↗")'
            assert hyperlink_formula == expected_formula, (
                f"Row {r_idx} in {sheet_name} ({file_path.name}) has invalid formula: {hyperlink_formula} (expected {expected_formula})"
            )
            valid_hyperlink_formula_count += 1

            # 3. Freshness must be one of the canonical states
            assert freshness in ("LIVE_ACTIVE", "APPROACHING_DEADLINE", "CLOSED"), (
                f"Row {r_idx} in {sheet_name} ({file_path.name}) has invalid freshness: {freshness}"
            )
            valid_freshness_count += 1

    assert total_data_rows == expected_rows, (
        f"{file_path.name} total rows {total_data_rows} != expected {expected_rows}"
    )
    assert valid_url_count == expected_rows, (
        f"{file_path.name} valid URLs {valid_url_count} != expected {expected_rows}"
    )
    assert valid_hyperlink_formula_count == expected_rows
    assert valid_freshness_count == expected_rows
