"""Unit tests for CLI subcommands: list-api-attributes and export-cluster-jobs.
Validates Decision JP-D04 in decision campaign cil-live-job-provenance-20261018.
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from cil_cli.main import main

PROJECT_ROOT = Path('/home/ubuntu/root/company-intelligence-platform')


def test_cli_list_api_attributes_cz_mpsv(capsys):
    ret = main(["list-api-attributes", "--source", "cz_mpsv", "--json-output"])
    assert ret == 0
    captured = capsys.readouterr()
    data = json.loads(captured.out)
    assert data["source"] == "cz_mpsv"
    assert data["provider_contract"] == "CZ_MPSV"
    assert data["attributes_count"] > 0
    # Check that portal_id is mapped to canonical job_id
    attrs = {a["attribute_name"]: a["mapped_canonical"] for a in data["attributes"]}
    assert attrs.get("portal_id") == "job_id"
    assert attrs.get("job_title_cz") == "job_title"


def test_cli_list_api_attributes_kr_alio(capsys):
    ret = main(["list-api-attributes", "--source", "kr_alio", "--json-output"])
    assert ret == 0
    captured = capsys.readouterr()
    data = json.loads(captured.out)
    assert data["source"] == "kr_alio"
    assert data["provider_contract"] == "JOB_ALIO"
    attrs = {a["attribute_name"]: a["mapped_canonical"] for a in data["attributes"]}
    assert attrs.get("institution_name") == "employer_name"
    assert attrs.get("posting_url") == "job_url"


def test_cli_list_api_attributes_us_usajobs(capsys):
    ret = main(["list-api-attributes", "--source", "us_usajobs", "--json-output"])
    assert ret == 0
    captured = capsys.readouterr()
    data = json.loads(captured.out)
    assert data["source"] == "us_usajobs"
    assert data["provider_contract"] == "US_USAJOBS"
    attrs = {a["attribute_name"]: a["mapped_canonical"] for a in data["attributes"]}
    assert attrs.get("PositionID") == "job_id"
    assert attrs.get("PositionURI") == "job_url"
    assert attrs.get("OrganizationName") == "employer_name"


def test_cli_export_cluster_jobs_by_signature(tmp_path, capsys):
    out_file = tmp_path / 'kr.xlsx'
    source = Path(__file__).parent / 'research/fixtures/selected-posting.json'
    assert main(['export-cluster-jobs', '--signature', 'KR:STATE_OWNED:ALL:REGULAR:ALL', '--snapshot', str(source), '--as-of', '2026-10-20T00:00:00Z', '--output', str(out_file), '--json-output']) == 0
    result = json.loads(capsys.readouterr().out)
    assert result['source_kind'] == 'fixture' and result['total_matched'] == 1
    workbook = openpyxl.load_workbook(out_file)
    assert 'Postings' in workbook.sheetnames and 'RawFields' in workbook.sheetnames


def test_cli_export_cluster_jobs_by_flags(tmp_path, capsys):
    source = Path(__file__).parent / 'research/fixtures/selected-posting.json'
    output = tmp_path / 'cz.xlsx'
    assert main(['export-cluster-jobs', '--country', 'CZ', '--sector', 'PUBLIC_INSTITUTION', '--snapshot', str(source), '--as-of', '2026-10-20T00:00:00Z', '--output', str(output), '--json-output']) == 0
    result = json.loads(capsys.readouterr().out)
    assert result['total_matched'] == 1
    workbook = openpyxl.load_workbook(output)
    assert any(cell.hyperlink for row in workbook['Postings'] for cell in row)


def test_cli_export_cluster_jobs_slovakia(tmp_path, capsys):
    source = Path(__file__).parent / 'research/fixtures/selected-posting.json'
    assert main(['export-cluster-jobs', '--country', 'SK', '--snapshot', str(source), '--as-of', '2026-10-20T00:00:00Z', '--output', str(tmp_path/'sk.xlsx'), '--json-output']) == 0
    result = json.loads(capsys.readouterr().out)
    assert result['total_matched'] == 1 and result['source_kind'] == 'fixture'
