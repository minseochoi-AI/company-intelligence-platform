from pathlib import Path
import openpyxl
import pytest
from cil_domain.taxonomy.lattice import DeterministicLatticeSynthesizer, DeterministicTaxonomyItem
from cil_reporting.xlsx.occupations import build_occupational_lattice_workbook, load_raw_isco_unit_groups


def test_load_raw_isco_unit_groups():
    with pytest.raises(ValueError, match='explicit taxonomy'):
        load_raw_isco_unit_groups()
    item = DeterministicTaxonomyItem('2512', 'Software developers', 'source description', '', '', [], '2', '25', '251')
    assert load_raw_isco_unit_groups([item]) == [{'isco_code': '2512', 'title_en': 'Software developers', 'description': 'source description'}]


def test_build_occupational_lattice_workbook(tmp_path: Path):
    target = tmp_path / 'occupations.xlsx'
    with pytest.raises(ValueError, match='explicit'):
        build_occupational_lattice_workbook(target)
    assert not target.exists()
    item = DeterministicTaxonomyItem('2512', 'Software developers', '', '', '', [], '2', '25', '251')
    graph = DeterministicLatticeSynthesizer(taxonomy={'2512': item}).build_deterministic_macro_graph()
    build_occupational_lattice_workbook(target, graph=graph)
    before = target.read_bytes()
    build_occupational_lattice_workbook(target, graph=graph)
    assert target.read_bytes() == before
    workbook = openpyxl.load_workbook(target)
    assert len(workbook.sheetnames) == 10
    assert workbook['01_Master_Taxonomy'].max_row == 2
    assert workbook['02_Career_Lattice_DAG'].max_row == 1
    assert workbook['07_Time_Series_Warehouse'].max_row == 1
    market = workbook['04_Market_Dynamics_ROI']
    headers = [c.value for c in market[1]]
    assert market.cell(2, headers.index('automation_exposure')+1).value == 'null'
